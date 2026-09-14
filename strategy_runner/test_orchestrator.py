import pathlib
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

import orchestrator


class StrategyRunnerTests(unittest.TestCase):
    def test_freqtrade_command_uses_path_on_windows_compatible_config(self):
        with patch.dict("os.environ", {}, clear=True), patch(
            "orchestrator.sys.executable", "/nonexistent/python"
        ), patch(
            "orchestrator.shutil.which", return_value=r"C:\\Miniconda3\\envs\\freqtrade313\\Scripts\\freqtrade.exe"
        ):
            self.assertEqual(
                [r"C:\\Miniconda3\\envs\\freqtrade313\\Scripts\\freqtrade.exe"],
                orchestrator.freqtrade_command({"freqtrade_bin": "freqtrade"}),
            )

    def test_freqtrade_command_honors_environment_override(self):
        with tempfile.TemporaryDirectory() as directory:
            executable = pathlib.Path(directory) / "freqtrade.exe"
            executable.touch()
            with patch.dict("os.environ", {"FREQTRADE_BIN": str(executable)}):
                self.assertEqual([str(executable)], orchestrator.freqtrade_command({}))

    def test_select_modes_supports_spot_futures_and_all(self):
        specs = [
            orchestrator.StrategySpec("Spot", "spot.py", "5m", "spot", False, False),
            orchestrator.StrategySpec("Futures", "futures.py", "15m", "futures", True, False),
        ]
        self.assertEqual(["Spot"], [item.name for item in orchestrator.select_modes(specs, "spot")])
        self.assertEqual(["Futures"], [item.name for item in orchestrator.select_modes(specs, "futures")])
        self.assertEqual(["Spot", "Futures"], [item.name for item in orchestrator.select_modes(specs, "all")])

    def test_discovers_expected_strategy_set(self):
        specs = orchestrator.discover_strategies()
        self.assertEqual(71, len(specs))
        self.assertIn("Bandtastic", {spec.name for spec in specs})
        self.assertTrue(any(spec.mode == "futures" for spec in specs))
        self.assertTrue(any(spec.lookahead_flag for spec in specs))
        self.assertTrue(all(spec.file.startswith(f"{spec.mode}/") for spec in specs))

    def test_scoring_disqualifies_low_trade_and_lookahead_rows(self):
        base = {
            "mode": "spot", "timerange": "20240101-20250101", "status": "success",
            "total_trades": 100, "profit_total": 0.1, "sharpe": 1.0, "sortino": 1.1,
            "calmar": 0.8, "max_drawdown_account": 0.2, "lookahead_flag": False,
        }
        rows = [
            {**base, "name": "Good"},
            {**base, "name": "FewTrades", "total_trades": 2},
            {**base, "name": "Biased", "lookahead_flag": True},
        ]
        result = {row["name"]: row for row in orchestrator.score_rows(rows, orchestrator.load_settings())}
        self.assertTrue(result["Good"]["eligible"])
        self.assertFalse(result["FewTrades"]["eligible"])
        self.assertFalse(result["Biased"]["eligible"])

    def test_writes_pandas_xlsx_report(self):
        rows = [{"name": "Example", "profit_total": 0.125, "rank": 1, "status": "success"}]
        with tempfile.TemporaryDirectory() as directory:
            path = orchestrator.write_xlsx_report(pathlib.Path(directory), rows)
            workbook = openpyxl.load_workbook(path)
            sheet = workbook["回测汇总"]
            self.assertEqual("A2", sheet.freeze_panes)
            self.assertEqual("A1:D2", sheet.auto_filter.ref)
            self.assertEqual("Example", sheet.cell(2, 1).value)
            self.assertEqual("0.00%", sheet.cell(2, 2).number_format)

    def test_xlsx_rejects_fields_without_chinese_labels(self):
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaisesRegex(ValueError, "unknown_metric"):
                orchestrator.write_xlsx_report(
                    pathlib.Path(directory), [{"name": "Example", "unknown_metric": 1}]
                )


if __name__ == "__main__":
    unittest.main()
